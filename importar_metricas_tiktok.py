# -*- coding: utf-8 -*-
"""Le o export do TikTok Studio e alimenta a metrica do ciclo semanal.

⚠️ POR QUE ISTO VALE MAIS QUE A PRINT DO PERFIL.

O ciclo semanal precisa saber quais posts mais viralizaram. As tres fontes
possiveis, medidas em 08/09/2026:

    view do Buffer   ZERO em posts de 6 e 7 dias que tinham 338, 359 e 142
                     views no TikTok. Nao e' atraso, e' ausencia — e onde ha'
                     numero ele nao guarda proporcao (0,06 a 0,71), entao nem
                     a ORDEM se preserva. NAO SERVE.
    print do perfil  numero real, mas lido a olho, so' o que cabe na tela, e
                     depende do Bryan mandar foto toda semana.
    export do Studio numero real, TODOS os posts, legivel por maquina.

O export ganha nos tres eixos. `studio.tiktok.com` -> Analytics -> periodo ->
exportar (CSV ou XLSX).

⚠️ O HISTORICO DO TIKTOK PARA EM 60 DIAS. O export nao e' arquivo permanente:
o que nao for importado dentro da janela some pra sempre. Por isso este
script SOMA ao que ja' existe em vez de substituir — ao contrario da print,
que e' um retrato e substitui.

⚠️ E O MAPA DE COLUNAS AINDA NAO FOI CONFERIDO CONTRA UM ARQUIVO REAL. Os
nomes mudam com o idioma da conta e com a versao do Studio. O script tenta
varios, e quando nao acha ele DIZ quais colunas encontrou, em vez de chutar
uma. Chutar coluna e' como se importa curtida achando que e' view.

Uso:
    python importar_metricas_tiktok.py --arquivo export.csv --canal modofuturo
    python importar_metricas_tiktok.py --arquivo export.xlsx --canal atefalhar --simular
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
import unicodedata
import zipfile
from datetime import datetime, timezone
from pathlib import Path

RAIZ = Path(__file__).resolve().parent
sys.path.insert(0, str(RAIZ))

from engine import melhores  # noqa: E402

# Nomes que a coluna pode ter. Ordem importa: o primeiro que casar vence.
#
# ⚠️ "views" ANTES de qualquer outra coisa que contenha "view". Se
# "video views" e "profile views" estiverem os dois no arquivo, pegar o
# errado troca o desempenho do POST pelo do PERFIL — e o ranking inteiro
# passa a medir outra coisa, sem dar sinal nenhum.
COLUNAS = {
    "titulo": ["video title", "titulo do video", "título do vídeo", "title",
               "titulo", "título", "video description", "descricao",
               "descrição", "post", "video"],
    "views": ["video views", "visualizacoes do video", "visualizações do vídeo",
              "views", "visualizacoes", "visualizações", "play count",
              "reproducoes", "reproduções"],
    "data": ["post time", "publish time", "data de publicacao",
             "data de publicação", "date", "data", "created time"],
}


def _norm(t: str) -> str:
    t = unicodedata.normalize("NFD", (t or "").strip().lower())
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", t)


def _achar_coluna(cabecalho: list[str], quais: list[str]) -> str | None:
    normal = {_norm(c): c for c in cabecalho}
    for alvo in quais:
        a = _norm(alvo)
        if a in normal:
            return normal[a]
    # segunda passada: casamento por conter, que e' mais frouxo
    for alvo in quais:
        a = _norm(alvo)
        for n, original in normal.items():
            if a in n:
                return original
    return None


def _numero(v) -> int | None:
    """"1.234", "1,234", "1 234" e "1.2K" viram inteiro."""
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    m = re.match(r"^([\d.,\s]+)\s*([kKmM])?$", s)
    if not m:
        return None
    corpo, suf = m.group(1).strip(), (m.group(2) or "").lower()

    # ⚠️ O MESMO PONTO SIGNIFICA COISAS OPOSTAS, e trocar as duas leituras
    # erra por 10x. MEDIDO em 08/09/2026: a primeira versao apagava todo `.`
    # e `,` antes de aplicar o sufixo, entao "1.2K" virava 12 * 1000 = 12000
    # em vez de 1200.
    #
    #     COM sufixo  ->  o separador e' DECIMAL   "1.2K"  = 1200
    #     SEM sufixo  ->  o separador e' MILHAR    "1.234" = 1234
    #
    # Um ranking com um valor inflado 10x poe o post errado em primeiro e
    # manda baixar 5 fontes no rumo dele. Nada na saida denunciaria.
    if suf:
        num = corpo.replace(" ", "").replace(",", ".")
        if num.count(".") > 1:
            return None
        try:
            valor = float(num)
        except ValueError:
            return None
        return int(round(valor * (1000 if suf == "k" else 1_000_000)))

    limpo = re.sub(r"[.,\s]", "", corpo)
    return int(limpo) if limpo.isdigit() else None


def _ler_content_ancorado(texto: str) -> list[dict]:
    """O export `Content.csv`, ancorado no LINK e nao na posicao da coluna.

    ⚠️ O EXPORT DO TIKTOK VEM QUEBRADO, e isso foi MEDIDO em 08/09/2026. Numa
    das 15 linhas do @achadinho.make a descricao continha virgulas e setas
    fora de aspas, e as colunas DESLOCARAM: o campo "Total views" veio com a
    URL do video dentro, e os numeros verdadeiros sobraram num campo extra:

        'Total views': 'https://www.tiktok.com/@achadinho.make/video/76820...'
        None:          "['5 de setembro', '34', '0', '0', '479']"

    Ler por nome de coluna nesse arquivo devolve lixo em silencio — e como o
    lixo e' uma string, `int()` estoura ou, pior, um numero errado passa.

    A ancora que NAO desloca e' o link: ele e' o unico campo com formato
    reconhecivel. Depois dele vem sempre, nesta ordem, post time, likes,
    comments, shares e views. O titulo e' tudo que esta' entre a data e o
    link, remontado.
    """
    linhas = list(csv.reader(io.StringIO(texto)))
    if len(linhas) < 2:
        return []
    saida = []
    for r in linhas[1:]:
        i = next((k for k, c in enumerate(r)
                  if str(c).startswith("https://www.tiktok.com/")), None)
        if i is None or len(r) < i + 6:
            continue
        v = _numero(r[i + 5])
        if v is None:
            continue
        saida.append({"titulo": " ".join(r[1:i]).strip(),
                      "views": v, "data": r[0], "link": r[i]})
    return saida


def ler(caminho: Path) -> tuple[list[dict], str]:
    """Devolve (linhas, aviso). Cada linha e' {titulo, views, data}."""
    # ⚠️ O Studio entrega ZIP. Abrir o zip aqui evita que cada canal vire um
    # passo manual de descompactar — e passo manual em rotina mensal e' passo
    # que um dia nao acontece.
    if caminho.suffix.lower() == ".zip":
        z = zipfile.ZipFile(caminho)
        nomes = [n for n in z.namelist() if n.lower().endswith(".csv")]
        if not nomes:
            return [], "zip sem csv dentro"
        texto = z.read(nomes[0]).decode("utf-8-sig", errors="replace")
        if "Video link" in texto.splitlines()[0]:
            linhas = _ler_content_ancorado(texto)
            if linhas:
                return linhas, (f"formato Content (ancorado no link), "
                                f"{len(linhas)} post(s)")
            return [], "Content.csv sem linha legivel"
        return [], ("este zip nao e' o export por POST. O Overview (uma linha "
                    "por DIA) se importa com importar_overview_tiktok.py")

    if caminho.suffix.lower() in (".xlsx", ".xls"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            return [], ("arquivo .xlsx e o openpyxl nao esta' instalado. "
                        "Exporte como CSV no Studio, ou: pip install openpyxl")
        wb = load_workbook(caminho, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        linhas = [[c for c in row] for row in ws.iter_rows(values_only=True)]
        if not linhas:
            return [], "planilha vazia"
        cab = [str(c) if c is not None else "" for c in linhas[0]]
        corpo = [dict(zip(cab, r)) for r in linhas[1:]]
    else:
        texto = caminho.read_text(encoding="utf-8-sig", errors="replace")
        # O Studio ja' exportou com ; em conta pt-BR; o sniff evita adivinhar.
        try:
            dial = csv.Sniffer().sniff(texto[:4000], delimiters=",;\t")
        except csv.Error:
            dial = csv.excel
        corpo = list(csv.DictReader(io.StringIO(texto), dialect=dial))
        cab = list(corpo[0].keys()) if corpo else []

    if not corpo:
        return [], "arquivo sem linhas de dados"

    c_tit = _achar_coluna(cab, COLUNAS["titulo"])
    c_vw = _achar_coluna(cab, COLUNAS["views"])
    c_dt = _achar_coluna(cab, COLUNAS["data"])
    if not c_tit or not c_vw:
        return [], (
            "nao achei as colunas de titulo e/ou views. NAO chutei nenhuma — "
            "chutar coluna e' como se importa curtida achando que e' view.\n"
            f"  colunas do arquivo: {', '.join(str(c) for c in cab)}\n"
            f"  titulo: {c_tit or 'NAO ACHEI'} | views: {c_vw or 'NAO ACHEI'}")

    saida = []
    for r in corpo:
        t = str(r.get(c_tit) or "").strip()
        v = _numero(r.get(c_vw))
        if t and v is not None:
            saida.append({"titulo": t, "views": v,
                          "data": str(r.get(c_dt) or "") if c_dt else ""})
    if not saida:
        return [], (f"colunas achadas ({c_tit} / {c_vw}) mas nenhuma linha "
                    f"com titulo e numero legiveis")
    return saida, f"colunas: titulo={c_tit!r}  views={c_vw!r}  data={c_dt!r}"


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--arquivo", required=True)
    p.add_argument("--canal", required=True)
    p.add_argument("--simular", action="store_true")
    a = p.parse_args()

    caminho = Path(a.arquivo)
    if not caminho.exists():
        sys.exit(f"nao achei {caminho}")
    linhas, aviso = ler(caminho)
    if not linhas:
        sys.exit(f"[!] {aviso}")
    print(aviso)
    linhas.sort(key=lambda x: -x["views"])
    print(f"\n{len(linhas)} post(s) lidos do export. Os 5 maiores:")
    for x in linhas[:5]:
        print(f"  {x['views']:>7}  {x['titulo'][:60].encode('ascii','replace').decode()}")

    if a.simular:
        print("\nSIMULADO — nada gravado.")
        return

    # ⚠️ SOMA ao que ja' existe, nao substitui: o historico do TikTok para em
    # 60 dias, e o que nao for importado na janela some pra sempre. A print
    # (um retrato) substitui; o export (um historico) acumula.
    try:
        atual = json.loads(melhores.MANUAL.read_text(encoding="utf-8"))
    except Exception:
        atual = {}
    ja = {x["titulo"]: x for x in atual.get(a.canal, {}).get("posts", [])}
    for x in linhas:
        anterior = ja.get(x["titulo"], {}).get("views", 0)
        # O maior vence: view so' sobe, e um export mais antigo nao pode
        # rebaixar um numero ja' conhecido.
        ja[x["titulo"]] = {"titulo": x["titulo"],
                           "views": max(int(anterior or 0), x["views"])}
    atual[a.canal] = {
        "medido_em": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "origem": f"export do TikTok Studio ({caminho.name})",
        "posts": sorted(ja.values(), key=lambda x: -x["views"]),
    }
    melhores.MANUAL.parent.mkdir(parents=True, exist_ok=True)
    melhores.MANUAL.write_text(json.dumps(atual, ensure_ascii=False, indent=1,
                                          sort_keys=True), encoding="utf-8")
    print(f"\n{len(ja)} post(s) no registro do {a.canal}.")
    top, fonte, _ = melhores.melhores(a.canal, 2)
    print(f"melhores da semana (fonte={fonte}):")
    for x in top:
        print(f"  {x['views']:>7}  {x['titulo'][:60].encode('ascii','replace').decode()}")


if __name__ == "__main__":
    main()
